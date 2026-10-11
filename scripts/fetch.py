"""Source adapters. Each returns {series_id: {period_key: value}} with period keys 'YYYY-MM' (monthly) or
'YYYY-MM-DD' (weekly, the week's reference date as the agency gives it). No API keys are needed.
Every adapter raises on failure; run.py records the failure and keeps the archived series."""
import csv
import datetime as dt
import io
import json
import re
import time
import urllib.parse
import urllib.request

UA = {"User-Agent": "promise-ledger/1.0 (public-interest data project; github.com/pantaleonfassbender-coder)"}
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
MON3 = {m[:3].upper(): i + 1 for i, m in enumerate(MONTHS)}


def get(url, data=None, headers=None, timeout=120, tries=4):
    h = dict(UA, **(headers or {}))
    for i in range(tries):
        try:
            req = urllib.request.Request(url, data=data, headers=h)
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.read()
        except Exception:
            if i == tries - 1:
                raise
            time.sleep(5 * (i + 1))


def mkey(y, m):
    return f"{y:04d}-{m:02d}"


# ---------------------------------------------------------------- BLS (public API v1, no key: 25 series x 10 years)
def bls(series_ids, first_year=2010):
    out = {s: {} for s in series_ids}
    this_year = dt.date.today().year
    for start in range(first_year, this_year + 1, 10):
        end = min(start + 9, this_year)
        for chunk in [series_ids[i:i + 25] for i in range(0, len(series_ids), 25)]:
            body = json.dumps({"seriesid": chunk, "startyear": str(start), "endyear": str(end)}).encode()
            d = json.loads(get("https://api.bls.gov/publicAPI/v1/timeseries/data/", data=body, headers={"Content-Type": "application/json"}))
            if d.get("status") != "REQUEST_SUCCEEDED":
                raise RuntimeError(f"BLS: {d.get('status')} {d.get('message')}")
            for s in d["Results"]["series"]:
                for o in s["data"]:
                    p = o["period"]
                    if not p.startswith("M") or p == "M13":
                        continue
                    try:
                        out[s["seriesID"]][mkey(int(o["year"]), int(p[1:]))] = float(o["value"])
                    except ValueError:
                        pass
    return out


# ---------------------------------------------------------------- EIA history spreadsheets (no key)
def eia(path, freq):
    import xlrd
    book = xlrd.open_workbook(file_contents=get("https://www.eia.gov/dnav/" + path))
    sh = book.sheet_by_name("Data 1")
    out = {}
    for r in range(3, sh.nrows):
        d, v = sh.cell_value(r, 0), sh.cell_value(r, 1)
        if v in ("", None) or d in ("", None):
            continue
        day = xlrd.xldate_as_datetime(d, book.datemode).date()
        out[mkey(day.year, day.month) if freq == "M" else day.isoformat()] = float(v)
    return {path: out}


# ---------------------------------------------------------------- Federal Reserve G.17 (seasonally adjusted text file)
def fed(codes):
    txt = get("https://www.federalreserve.gov/releases/g17/ipdisk/ip_sa.txt").decode("latin-1")
    out = {c: {} for c in codes}
    for line in txt.splitlines():
        m = re.match(r'^"([A-Z0-9_.]+)"\s+(\d{4})\s+(.*)$', line)
        if not m or m.group(1) not in out:
            continue
        year = int(m.group(2))
        for i, v in enumerate(m.group(3).split()[:12]):
            try:
                out[m.group(1)][mkey(year, i + 1)] = float(v)
            except ValueError:
                pass
    return out


# ---------------------------------------------------------------- Treasury fiscaldata (no key)
FD = "https://api.fiscaldata.treasury.gov/services/api/fiscal_service/"


def fd_all(endpoint, params):
    rows, page = [], 1
    while True:
        q = dict(params, **{"page[size]": "1000", "page[number]": str(page)})
        d = json.loads(get(FD + endpoint + "?" + urllib.parse.urlencode(q, safe=":,[]()")))
        rows += d["data"]
        if page >= d["meta"]["total-pages"]:
            return rows
        page += 1


def mts_customs():
    rows = fd_all("v1/accounting/mts/mts_table_4", {"filter": "classification_desc:eq:Customs Duties,record_date:gte:2014-10-01", "sort": "record_date"})
    gross, net = {}, {}
    for r in rows:
        k = r["record_date"][:7]
        if k in gross:
            continue
        gross[k] = float(r["current_month_gross_rcpt_amt"]) / 1e6
        net[k] = float(r["current_month_net_rcpt_amt"]) / 1e6
    return {"gross": gross, "net": net}


def mts_outlays():
    """Monthly outlays and deficit from the fiscal-year blocks of Table 1 (each release repeats the prior and the
    current fiscal year; later releases override earlier ones)."""
    latest = fd_all("v1/accounting/mts/mts_table_1", {"sort": "-record_date", "page[size]": "1", "fields": "record_date"})
    dates = [f"{y}-09-30" for y in range(2015, int(latest[0]["record_date"][:4]) + 1)] + [latest[0]["record_date"]]
    outlays, deficit = {}, {}
    for d in sorted(set(dates)):
        rows = fd_all("v1/accounting/mts/mts_table_1", {"filter": f"record_date:eq:{d}"})
        rows.sort(key=lambda r: int(r["src_line_nbr"]))
        fy = None
        for r in rows:
            m = re.match(r"FY (\d{4})", r["classification_desc"])
            if m:
                fy = int(m.group(1))
                continue
            if r["record_type_cd"] != "MTH" or fy is None or r["classification_desc"] not in MONTHS:
                continue
            mo = MONTHS.index(r["classification_desc"]) + 1
            k = mkey(fy - 1 if mo >= 10 else fy, mo)
            if r["current_month_gross_outly_amt"] not in (None, "null"):
                outlays[k] = float(r["current_month_gross_outly_amt"]) / 1e6
                deficit[k] = float(r["current_month_dfct_sur_amt"]) / 1e6
    return {"outlays": outlays, "deficit": deficit}


def dts_customs():
    """Customs deposits (Treasury General Account), summed by week ending Friday. The line was renamed on
    2025-11-10; both names are read."""
    daily = {}
    for label in ("DHS - Customs and Certain Excise Taxes", "DHS - Customs Duties, Taxes, and Fees"):
        rows = fd_all("v1/accounting/dts/deposits_withdrawals_operating_cash",
                      {"filter": f"transaction_catg:eq:{label},transaction_type:eq:Deposits,record_date:gte:2022-10-01", "fields": "record_date,transaction_today_amt"})
        for r in rows:
            daily[r["record_date"]] = daily.get(r["record_date"], 0.0) + float(r["transaction_today_amt"])
    weeks = {}
    for d, v in daily.items():
        day = dt.date.fromisoformat(d)
        fri = day + dt.timedelta(days=(4 - day.weekday()) % 7)
        weeks[fri.isoformat()] = weeks.get(fri.isoformat(), 0.0) + v
    last = max(daily)            # drop the current, incomplete week
    lastfri = dt.date.fromisoformat(last) + dt.timedelta(days=(4 - dt.date.fromisoformat(last).weekday()) % 7)
    if lastfri.isoformat() != last:
        weeks.pop(lastfri.isoformat(), None)
    return {"customs": weeks}


# ---------------------------------------------------------------- Census FT-900, Exhibit 1 history since 1992 (seasonally adjusted)
def ft900():
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(get("https://www.census.gov/foreign-trade/statistics/historical/exhibit_history.xlsx")), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    year, out = None, {}
    for r in ws.iter_rows(values_only=True):
        first = str(r[0]).strip() if r[0] is not None else ""
        if first.startswith("Part B"):
            break
        if re.fullmatch(r"\d{4}", first):
            year = int(first)
            continue
        name = re.sub(r"\s*\(R\)\s*", "", first)
        if year and name in MONTHS and isinstance(r[2], (int, float)):
            out[mkey(year, MONTHS.index(name) + 1)] = float(r[2])
    if not out:
        raise RuntimeError("FT-900: no monthly rows found")
    return {"goods": out}


# ---------------------------------------------------------------- CBP nationwide encounters
# cbp.gov and ice.gov refuse automated requests (HTTP 403). The ledger does not disguise itself as a browser; it reads
# the agencies' own published files from the Deportation Data Project's public CC0 mirrors on GitHub instead.
def gh_tree(repo):
    return json.loads(get(f"https://api.github.com/repos/deportationdata/{repo}/git/trees/main?recursive=1"))["tree"]


def gh_raw(repo, path, lfs=False):
    base = "https://media.githubusercontent.com/media" if lfs else "https://raw.githubusercontent.com"
    return get(f"{base}/deportationdata/{repo}/main/" + urllib.parse.quote(path), timeout=600)


def cbp():
    """CBP's Nationwide Encounters files (each covers four fiscal years); newer files override older months."""
    paths = sorted(t["path"] for t in gh_tree("cbp") if re.search(r"encounters/validation/nationwide-encounters-fy\d\d-fy\d\d.*-aor\.csv$", t["path"]))
    if not paths:
        raise RuntimeError("CBP mirror: no Nationwide Encounters file found")
    out = {}
    for path in paths:                       # sorted by covered years: later files win
        per = {}
        for r in csv.DictReader(io.StringIO(gh_raw("cbp", path).decode("utf-8-sig", "replace"))):
            if r["Component"].strip() != "U.S. Border Patrol" or r["Land Border Region"].strip() != "Southwest Land Border" or r["Encounter Type"].strip() != "Apprehensions":
                continue
            fy, mo = int(r["Fiscal Year"].strip()[:4]), MON3[r["Month (abbv)"].strip().upper()[:3]]
            k = mkey(fy - 1 if mo >= 10 else fy, mo)
            per[k] = per.get(k, 0.0) + float(r["Encounter Count"].replace(",", ""))
        out.update(per)
    return {"usbp_sw": out, "_file": paths[-1]}


# ---------------------------------------------------------------- ICE detention statistics (average daily population)
def ice_adp():
    """Monthly average daily population in ICE detention (all agencies, all criminality classes), compiled from
    ICE's published detention statistics files by the Deportation Data Project (CC0)."""
    import openpyxl
    path = "data/adp-and-stay-length-by-agency-monthly.xlsx"
    ws = openpyxl.load_workbook(io.BytesIO(gh_raw("ice-detention-management", path, lfs=True)), read_only=True, data_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    hdr = list(next(rows))
    ix = {h: i for i, h in enumerate(hdr)}
    out, pulled = {}, None
    for r in rows:
        if r[ix["agency"]] != "Total" or r[ix["criminality"]] != "Total" or r[ix["adp"]] in (None, 0):
            continue
        d = r[ix["date"]]
        out[mkey(d.year, d.month)] = float(r[ix["adp"]])
        p = r[ix["pull_date"]]
        if p and (pulled is None or p > pulled):
            pulled = p
    if not out:
        raise RuntimeError("ICE mirror: no monthly totals found")
    return {"total": out, "_data_as_of": pulled.date().isoformat() if pulled else None}


# ---------------------------------------------------------------- ICE removals released under FOIA (Deportation Data Project)
GH = "https://api.github.com/repos/deportationdata/ice"


def foia_removals(cache, budget_s=1500):
    """Monthly removals and removals with a criminal conviction, from ICE's removal records (inputs/*REM*.xlsx,
    Git LFS). A file is parsed again only when its content hash changes; aggregates are cached."""
    tree = gh_tree("ice")
    files = sorted(t["path"] for t in tree if re.search(r"inputs/.*_REM_FY20\d\d.*\.xlsx$", t["path"]))
    t0 = time.time()
    for path in files:
        pointer = get("https://raw.githubusercontent.com/deportationdata/ice/main/" + urllib.parse.quote(path)).decode("utf-8", "replace")
        oid = re.search(r"oid sha256:([0-9a-f]+)", pointer).group(1)
        if cache.get(path, {}).get("oid") == oid:
            continue
        if time.time() - t0 > budget_s:
            break                                   # the next run continues; the cache keeps what is done
        import xlsxstream
        data = gh_raw("ice", path, lfs=True)
        through, hdr, allm, conv = None, None, {}, {}
        for r in xlsxstream.rows(data):
            if hdr is None:
                for v in r.values():
                    if isinstance(v, str):
                        m = re.search(r"\((\d\d/\d\d/\d{4}) - (\d\d/\d\d/\d{4})\)", v)
                        if m:
                            through = dt.datetime.strptime(m.group(2), "%m/%d/%Y").date().isoformat()
                if "Departed Date" in r.values():
                    hdr = {v: k for k, v in r.items() if isinstance(v, str)}
                continue
            d = xlsxstream.to_date(r.get(hdr["Departed Date"]))
            if d is None:
                continue
            k = mkey(d.year, d.month)
            allm[k] = allm.get(k, 0) + 1
            if "Convicted" in str(r.get(hdr.get("Case Criminality"), "")):
                conv[k] = conv.get(k, 0) + 1
        cache[path] = dict(oid=oid, through=through, all=allm, convicted=conv)
    allm, conv, through = {}, {}, None
    for path, c in cache.items():
        for k, v in c["all"].items():
            allm[k] = allm.get(k, 0) + v
        for k, v in c["convicted"].items():
            conv[k] = conv.get(k, 0) + v
        if c.get("through") and (through is None or c["through"] > through):
            through = c["through"]
    if through:                                   # the last month is complete only if the data run to its end
        last = dt.date.fromisoformat(through)
        if (last + dt.timedelta(days=1)).month == last.month:
            for d in (allm, conv):
                d.pop(mkey(last.year, last.month), None)
    return {"all": allm, "convicted": conv, "_through": through}
